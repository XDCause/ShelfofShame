"""Import the hobby supplies spreadsheet into PostgreSQL.

Usage (from the project folder):
    .venv\\Scripts\\python.exe scripts\\import_spreadsheet.py "C:\\path\\to\\Hobby_Supplies_Inventory.xlsx"

Reads two sheets:
    Lists    -> categories and brands (added if missing)
    Supplies -> supplies (inserted, or updated if that brand + product code already exists)

Safe to re-run: rows are matched on (brand, product code), so nothing is duplicated.
Everything happens in one transaction: if any row fails, nothing is written.
"""
import sys
import warnings
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import psycopg
from openpyxl import load_workbook

# openpyxl can't read Excel's newer dropdown (data validation) format. Harmless here:
# we only read values and never save the workbook.
warnings.filterwarnings("ignore", message="Data Validation extension")

PROJECT_DIR = Path(__file__).resolve().parent.parent
HEADER_ROW = 4  # Supplies sheet: title in row 1, notes in row 2, headers in row 4

# Spreadsheet header -> key used below (and mostly the supplies column name)
SUPPLIES_COLUMNS = {
    "Item ID": "item_ref",
    "Brand": "brand",
    "Product Code": "product_code",
    "Item / Colour Name": "name",
    "Category": "category",
    "Paint Type": "paint_type",
    "Finish": "finish",
    "Size": "size",
    "Qty On Hand": "qty_on_hand",
    "Reorder At": "reorder_at",
    "Condition": "condition",
    "Unit Price ($)": "unit_price",
    "Purchase Date": "purchase_date",
    "Purchased From": "purchased_from",
    "Storage": "storage",
    "Notes": "notes",
}
# Stock Status and Total Value are spreadsheet formulas; the v_supplies view calculates them.

# Corrections agreed 2026-10-11: these paints had Finish 'N/A' in the spreadsheet.
FINISH_FIXES = {
    ("Tamiya", "T81764"): "Flat / Matt",  # XF-64 Red Brown
    ("Tamiya", "T81776"): "Flat / Matt",  # XF-76 Grey Green
    ("Tamiya", "T81771"): "Flat / Matt",  # XF-71 Cockpit Green
    ("Tamiya", "T81510"): "Metallic",     # X-10 Gun Metal
}


def load_env():
    """Read KEY=VALUE pairs from the project's .env file."""
    env = {}
    for line in (PROJECT_DIR / ".env").read_text().splitlines():
        line = line.strip()
        if line and not line.startswith("#") and "=" in line:
            key, value = line.split("=", 1)
            env[key.strip()] = value.strip()
    return env


def clean(value):
    """Trim text and turn empty cells into None (NULL in the database)."""
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def read_lists(wb):
    """Lists sheet: header in row 1, values below. Returns {header: [values]}."""
    ws = wb["Lists"]
    headers = [clean(c.value) for c in ws[1]]
    lists = {h: [] for h in headers if h}
    for row in ws.iter_rows(min_row=2, max_row=40, values_only=True):
        for header, value in zip(headers, row):
            value = clean(value)
            if header and value is not None:
                lists[header].append(value)
    return lists


def read_supplies(wb):
    """Supplies sheet: one dict per row that has a Brand or a Name."""
    ws = wb["Supplies"]
    headers = [clean(c.value) for c in ws[HEADER_ROW]]
    missing = [h for h in SUPPLIES_COLUMNS if h not in headers]
    if missing:
        sys.exit(f"Supplies sheet is missing columns: {missing}")
    index = {SUPPLIES_COLUMNS[h]: headers.index(h) for h in SUPPLIES_COLUMNS}

    rows = []
    for excel_row, values in enumerate(ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True),
                                       start=HEADER_ROW + 1):
        row = {key: clean(values[i]) for key, i in index.items()}
        if row["brand"] is None and row["name"] is None:
            continue  # empty row (only the pre-filled formulas)
        row["excel_row"] = excel_row
        # Keep the spreadsheet's numbering: SUP-0007 -> supply_id 7
        ref = row.pop("item_ref")
        row["supply_id"] = int(ref.split("-")[1]) if ref else excel_row - HEADER_ROW
        if isinstance(row["purchase_date"], datetime):
            row["purchase_date"] = row["purchase_date"].date()
        if row["unit_price"] is not None:
            # Excel stores 3.95 as a binary float; Decimal keeps exact cents for NUMERIC(10,2)
            row["unit_price"] = round(Decimal(str(row["unit_price"])), 2)
        fix = FINISH_FIXES.get((row["brand"], row["product_code"]))
        if fix:
            row["finish"] = fix
        rows.append(row)
    return rows


def validate(rows, lists, storage_names):
    """Check every value before touching the database, so errors name the Excel row."""
    allowed = {
        "category": set(lists["Category"]),
        "brand": set(lists["Brand"]),
        "paint_type": set(lists["Paint Type"]),
        "finish": set(lists["Finish"]),
        "condition": set(lists["Condition"]),
        "storage": storage_names,
    }
    errors = []
    for row in rows:
        r = row["excel_row"]
        if row["name"] is None:
            errors.append(f"row {r}: no Item / Colour Name")
        if row["category"] is None:
            errors.append(f"row {r}: no Category")
        for key, values in allowed.items():
            if row[key] is not None and row[key] not in values:
                errors.append(f"row {r}: {key} '{row[key]}' is not in the allowed list")
        for key in ("qty_on_hand", "reorder_at"):
            if row[key] is not None and (not isinstance(row[key], int) or row[key] < 0):
                errors.append(f"row {r}: {key} must be a whole number >= 0, got {row[key]!r}")
        if row["purchase_date"] is not None and not isinstance(row["purchase_date"], date):
            errors.append(f"row {r}: purchase date {row['purchase_date']!r} is not a date")
    return errors


UPSERT_SQL = """
INSERT INTO supplies (
    supply_id, brand_id, product_code, name, category_id, paint_type, finish, size,
    qty_on_hand, reorder_at, condition, unit_price, purchase_date, purchased_from,
    location_id, notes)
VALUES (
    %(supply_id)s, %(brand_id)s, %(product_code)s, %(name)s, %(category_id)s, %(paint_type)s,
    %(finish)s, %(size)s, COALESCE(%(qty_on_hand)s, 0), %(reorder_at)s, %(condition)s,
    %(unit_price)s, %(purchase_date)s, %(purchased_from)s, %(location_id)s, %(notes)s)
ON CONFLICT (brand_id, product_code) DO UPDATE SET
    name           = EXCLUDED.name,
    category_id    = EXCLUDED.category_id,
    paint_type     = EXCLUDED.paint_type,
    finish         = EXCLUDED.finish,
    size           = EXCLUDED.size,
    qty_on_hand    = EXCLUDED.qty_on_hand,
    reorder_at     = EXCLUDED.reorder_at,
    condition      = EXCLUDED.condition,
    unit_price     = EXCLUDED.unit_price,
    purchase_date  = EXCLUDED.purchase_date,
    purchased_from = EXCLUDED.purchased_from,
    location_id    = EXCLUDED.location_id,
    notes          = EXCLUDED.notes
RETURNING (xmax = 0) AS inserted  -- xmax is 0 for a brand-new row, non-zero for an updated one
"""


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    xlsx = Path(sys.argv[1])

    # data_only=True gives the values Excel last calculated, not the formula text
    wb = load_workbook(xlsx, data_only=True, read_only=True)
    lists = read_lists(wb)
    rows = read_supplies(wb)
    print(f"Read {len(rows)} supplies from {xlsx.name}")

    env = load_env()
    conninfo = dict(host="127.0.0.1", port=env.get("POSTGRES_PORT", "5433"),
                    dbname=env["POSTGRES_DB"], user=env["POSTGRES_USER"],
                    password=env["POSTGRES_PASSWORD"])

    # 'with connect' = one transaction: commit if the block finishes, roll back on any error
    with psycopg.connect(**conninfo) as conn, conn.cursor() as cur:
        # Label the quantity changes for the stock history trigger (06_stock_history.sql).
        cur.execute("SELECT set_config('app.stock_reason', 'import', true)")
        cur.execute("SELECT name, location_id FROM storage_locations")
        locations = dict(cur.fetchall())

        errors = validate(rows, lists, set(locations))
        if errors:
            print("Nothing imported. Fix these in the spreadsheet first:")
            print("\n".join("  " + e for e in errors))
            sys.exit(1)

        for table, key in (("categories", "Category"), ("brands", "Brand")):
            for name in lists[key]:
                cur.execute(f"INSERT INTO {table} (name) VALUES (%s) ON CONFLICT (name) DO NOTHING",
                            (name,))
                if cur.rowcount:
                    print(f"  added {table[:-1] if table == 'brands' else 'category'}: {name}")

        cur.execute("SELECT name, category_id FROM categories")
        categories = dict(cur.fetchall())
        cur.execute("SELECT name, brand_id FROM brands")
        brands = dict(cur.fetchall())

        inserted = updated = 0
        for row in rows:
            params = dict(row,
                          brand_id=brands.get(row["brand"]),
                          category_id=categories[row["category"]],
                          location_id=locations.get(row["storage"]))
            cur.execute(UPSERT_SQL, params)
            if cur.fetchone()[0]:
                inserted += 1
            else:
                updated += 1

        # We supplied supply_id ourselves, so move the SERIAL counter past the highest one;
        # otherwise the next new item (e.g. from the API) would try to reuse id 1.
        cur.execute("SELECT setval(pg_get_serial_sequence('supplies', 'supply_id'),"
                    " COALESCE(MAX(supply_id), 1)) FROM supplies")

    print(f"Done: {inserted} inserted, {updated} updated.")


if __name__ == "__main__":
    main()
